"""Assemble final/: the files the team sends out, rebuilt from the current deliverables.

Run from the repository root after build_docx.py:
.venv/bin/python src/report/make_final.py
  → final/Elektrozavod_report.docx           the report (title page, sections 1–7, appendices А–Ж)
    final/Elektrozavod_report.pdf            the same rendered by LibreOffice (Times New Roman must be visible to it)
    final/Elektrozavod_calendar.pdf          the campaign calendar as a separate page
    final/Elektrozavod_budget_calendar.xlsx  budget variants А and Б and the calendar, for editing
Old files in final/ are replaced; nothing else is kept there.
"""
import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from build_appendices import MONTHS, TAXES, basis

ROOT = Path(__file__).resolve().parents[2]
FINAL = ROOT / 'final'
FONT = 'Times New Roman'


def sheet(wb, title, header, rows, widths, total_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if total_col:
        n = ws.max_row
        col = get_column_letter(total_col)
        ws.append([''] * (total_col - 2) + ['Итого', f'=SUM({col}2:{col}{n})'])
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT, bold=True)
        for row in ws.iter_rows(min_row=2, min_col=total_col - 3, max_col=total_col):
            for c in row:
                if isinstance(c.value, (int, float)) or str(c.value).startswith('='):
                    c.number_format = '# ##0'
    ws.freeze_panes = 'A2'
    return ws


def budget_rows(path):
    b = pd.read_csv(path)
    b = b.sort_values('phase', key=lambda s: s.map(list(MONTHS).index), kind='stable')
    rows = []
    for r in b.itertuples():
        month = 'ноябрь–март' if r.category == 'Резерв' else MONTHS[r.phase]
        price = None if pd.isna(r.unit_cost_rub) else r.unit_cost_rub
        qty = None if r.category in TAXES else r.qty
        rows.append([r.line_id, month, r.category, r.item.removeprefix('Вариант Б: '), '' if r.category in TAXES
                     else r.unit, price, qty, int(r.total_rub), basis(r.cost_source, r.category)])
    return rows


def workbook(out):
    wb = Workbook()
    wb.remove(wb.active)
    header = ['№', 'Месяц', 'Категория', 'Статья расходов', 'Единица', 'Цена, ₽ без НДС', 'Кол-во', 'Сумма, ₽',
              'Основание цены']
    widths = [7, 14, 26, 60, 16, 14, 9, 14, 28]
    sheet(wb, 'Вариант А', header, budget_rows(ROOT / 'deliverables/budget.csv'), widths, total_col=8)
    sheet(wb, 'Вариант Б', header, budget_rows(ROOT / 'deliverables/budget_variant_b.csv'), widths, total_col=8)
    r = pd.read_csv(ROOT / 'deliverables/roadmap.csv')
    rows = []
    for x in r.itertuples():
        code, what = x.activity.split(':', 1)
        rows.append([code, pd.Timestamp(x.week_start).strftime('%d.%m.%Y'), what.strip().removeprefix('Вариант Б: '),
                     x.audience, x.channel, x.owner_role, x.deliverable, x.kpi,
                     'да' if x.status == 'variant_b' else ''])
    sheet(wb, 'Календарь', ['№', 'Неделя', 'Что делаем', 'Аудитория', 'Канал', 'Кто отвечает', 'Результат',
                            'Показатель', 'Только в варианте Б'], rows, [6, 12, 55, 30, 30, 30, 35, 45, 12])
    wb.save(out)


def to_pdf(docx, out):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(docx)],
                       check=True, capture_output=True)
        shutil.move(Path(tmp) / (docx.stem + '.pdf'), out)


def main():
    argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter).parse_args()
    if FINAL.exists():
        shutil.rmtree(FINAL)
    FINAL.mkdir()
    report = FINAL / 'Elektrozavod_report.docx'
    shutil.copy2(ROOT / 'deliverables/Elektrozavod_report.docx', report)
    to_pdf(report, FINAL / 'Elektrozavod_report.pdf')
    shutil.copy2(ROOT / 'deliverables/appendix_timeline.pdf', FINAL / 'Elektrozavod_calendar.pdf')
    workbook(FINAL / 'Elektrozavod_budget_calendar.xlsx')
    for f in sorted(FINAL.iterdir()):
        print(f'{f.relative_to(ROOT)}  {f.stat().st_size // 1024} KB')


if __name__ == '__main__':
    main()
